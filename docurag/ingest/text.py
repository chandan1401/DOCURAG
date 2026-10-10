"""DOCX and plaintext loaders that emit PageData."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Iterable, List, Sequence
from xml.etree import ElementTree

from PIL import Image, ImageDraw

from docurag.config import get_settings
from docurag.ingest.tables import table_to_markdown
from docurag.schemas import PageData

WORDS_PER_PAGE = 500


def _write_page_preview(document_id: str, page_number: int, title: str = "") -> str:
    settings = get_settings()
    settings.ensure_dirs()
    out_path = settings.pages_dir / f"{document_id}_p{page_number}.png"
    img = Image.new("RGB", (612, 792), "white")
    draw = ImageDraw.Draw(img)
    draw.text((48, 48), title[:80] or f"{document_id} page {page_number}", fill=(32, 32, 32))
    img.save(out_path)
    return str(out_path)


def _paginate_text(text: str, document_id: str) -> List[PageData]:
    words = re.findall(r"\S+", text)
    if not words:
        return [
            PageData(
                page=1,
                text="",
                page_image_path=_write_page_preview(document_id, 1, document_id),
            )
        ]
    pages: List[PageData] = []
    for start in range(0, len(words), WORDS_PER_PAGE):
        page_number = (start // WORDS_PER_PAGE) + 1
        page_text = " ".join(words[start : start + WORDS_PER_PAGE])
        pages.append(
            PageData(
                page=page_number,
                text=page_text,
                ocr_used=False,
                tables_md=[],
                figures=[],
                page_image_path=_write_page_preview(document_id, page_number, page_text),
            )
        )
    return pages


def load_txt(path: str | Path, document_id: str) -> List[PageData]:
    """Load a plaintext document with encoding fallback and virtual pages."""
    path_obj = Path(path)
    data = path_obj.read_bytes()
    text = ""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    text = text.replace("\f", "\n\n")
    return _paginate_text(text, document_id)


def _extract_docx_text_with_python_docx(path: Path) -> tuple[str, List[str]]:
    from docx import Document

    doc = Document(str(path))
    parts: List[str] = []
    tables: List[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = getattr(para.style, "name", "") or ""
        if style.startswith("Heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "1"
            parts.append(f"{'#' * min(int(level), 6)} {text}")
        else:
            parts.append(text)
    for table in doc.tables:
        raw = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        md = table_to_markdown(raw)
        if md:
            tables.append(md)
    return "\n\n".join(parts), tables


def _extract_docx_text_from_zip(path: Path) -> tuple[str, List[str]]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    paras: List[str] = []
    for para in root.findall(".//w:p", ns):
        text = "".join(node.text or "" for node in para.findall(".//w:t", ns)).strip()
        if text:
            paras.append(text)
    return "\n\n".join(paras), []


def load_docx(path: str | Path, document_id: str) -> List[PageData]:
    """Load DOCX paragraphs and tables with a zip/XML fallback."""
    path_obj = Path(path)
    try:
        text, tables = _extract_docx_text_with_python_docx(path_obj)
    except Exception:
        text, tables = _extract_docx_text_from_zip(path_obj)

    pages = _paginate_text(text, document_id)
    if tables:
        pages[0].tables_md.extend(tables)
    return pages

